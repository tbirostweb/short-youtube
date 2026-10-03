"""Génère un modèle Excel (.xlsx) à remplir pour le mode batch.

Usage : python main.py template mes_shorts.xlsx
Puis   : python main.py batch mes_shorts.xlsx
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"

# (clé interne, en-tête affiché, largeur, aide courte)
COLUMNS = [
    ("channel", "channel", 10, "Chaîne cible (fr1, fr2, en1, en2)"),
    ("titre", "titre", 34, "Titre YouTube (utilisé tel quel)"),
    ("script", "script", 60, "Texte lu par la voix off. Vide = généré par IA si clé."),
    ("topic", "topic", 30, "Thème (si pas de script, sert à générer/compléter)"),
    ("keywords", "keywords", 26, "Mots-clés (virgule) -> b-roll Pexels + tags"),
    ("hashtags", "hashtags", 24, "Hashtags (virgule)"),
    ("cta", "cta", 26, "Phrase de fin (call-to-action)"),
    ("voix", "voix", 24, "Voix edge-tts (vide = voix de la chaîne)"),
    ("musique", "musique", 18, "Fichier dans data/assets/music, ou 'aucune'"),
    ("fond", "fond", 12, "auto | pexels | local | color"),
    ("categorie", "categorie", 11, "categoryId YouTube (nombre). Vide = chaîne."),
    ("traduire_vers", "traduire_vers", 15, "Autre chaîne : duplique la version traduite"),
    ("publier", "publier", 10, "oui / non"),
    ("date_publication", "date_publication", 20, "ex 2026-08-01 09:00 (vide = maintenant)"),
]

EXAMPLES = [
    {
        "channel": "fr1",
        "titre": "3 faits fous sur les trous noirs",
        "script": ("Un trou noir peut littéralement ralentir le temps. Approche-toi "
                   "assez près et une seconde pour toi dure des années ailleurs. "
                   "Deuxième fait : rien ne s'en échappe, pas même la lumière. Et le "
                   "plus dingue, au centre de presque chaque galaxie s'en cache un, "
                   "géant. Abonne-toi pour un fait par jour."),
        "topic": "faits sur les trous noirs",
        "keywords": "space, galaxy, black hole, stars",
        "hashtags": "#espace, #science, #shorts",
        "cta": "Abonne-toi pour un fait par jour.",
        "voix": "",
        "musique": "",
        "fond": "pexels",
        "categorie": "",
        "traduire_vers": "en1",
        "publier": "non",
        "date_publication": "",
    },
    {
        "channel": "fr2",
        "titre": "La règle des 2 minutes pour arrêter de procrastiner",
        "script": "",
        "topic": "arrêter de procrastiner avec la règle des deux minutes",
        "keywords": "focus, productivity, work, clock",
        "hashtags": "#productivité, #motivation, #shorts",
        "cta": "Abonne-toi pour un conseil par jour.",
        "voix": "fr-FR-HenriNeural",
        "musique": "aucune",
        "fond": "color",
        "categorie": "27",
        "traduire_vers": "",
        "publier": "non",
        "date_publication": "2026-08-01 09:00",
    },
]

VOICES_FR = [
    "fr-FR-DeniseNeural (F)", "fr-FR-HenriNeural (H)", "fr-FR-EloiseNeural (F)",
    "fr-FR-VivienneMultilingualNeural (F)", "fr-FR-RemyMultilingualNeural (H)",
]
VOICES_EN = [
    "en-US-AriaNeural (F)", "en-US-GuyNeural (H)", "en-US-JennyNeural (F)",
    "en-GB-SoniaNeural (F)", "en-GB-RyanNeural (H)",
]
CATEGORIES = [
    ("1", "Film & Animation"), ("10", "Music"), ("17", "Sports"),
    ("20", "Gaming"), ("22", "People & Blogs"), ("23", "Comedy"),
    ("24", "Entertainment"), ("25", "News & Politics"), ("26", "Howto & Style"),
    ("27", "Education"), ("28", "Science & Technology"),
]


def _style_header(cell) -> None:
    cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
    cell.fill = PatternFill("solid", fgColor="0E1116")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def make_template(output: str | Path = "shorts_template.xlsx") -> Path:
    out = Path(output)
    wb = Workbook()

    # ============ Feuille "Shorts" ============
    ws = wb.active
    ws.title = "Shorts"
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c, (_key, header, width, _help) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=header)
        _style_header(cell)
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = width

    for r, ex in enumerate(EXAMPLES, start=2):
        for c, (key, *_rest) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=ex.get(key, ""))
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=key in {"script", "topic"})
            cell.border = border

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30

    # --- menus déroulants (aides à la saisie, non bloquants) ---
    def add_list(col_key: str, options: str) -> None:
        idx = next(i for i, col in enumerate(COLUMNS, start=1) if col[0] == col_key)
        letter = get_column_letter(idx)
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        dv.error = None
        dv.showErrorMessage = False  # simple aide, on n'empêche pas d'autres valeurs
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}1000")

    add_list("channel", "fr1,fr2,en1,en2")
    add_list("fond", "auto,pexels,local,color")
    add_list("publier", "oui,non")
    add_list("traduire_vers", "en1,en2,fr1,fr2")

    # ============ Feuille "Notice" ============
    nz = wb.create_sheet("Notice")
    nz.column_dimensions["A"].width = 20
    nz.column_dimensions["B"].width = 85
    title = nz.cell(row=1, column=1, value="Comment remplir ce fichier")
    title.font = Font(name=FONT, bold=True, size=14)
    nz.cell(row=2, column=1,
            value="1 ligne = 1 Short. Seul le nécessaire pour un script est obligatoire "
                  "(soit la colonne 'script', soit 'topic' + une clé LLM dans .env).")\
        .font = Font(name=FONT, size=10, italic=True)

    hdr_a = nz.cell(row=4, column=1, value="Colonne")
    hdr_b = nz.cell(row=4, column=2, value="À quoi ça sert")
    for cell in (hdr_a, hdr_b):
        _style_header(cell)
    for i, (_key, header, _w, help_txt) in enumerate(COLUMNS, start=5):
        nz.cell(row=i, column=1, value=header).font = Font(name=FONT, bold=True, size=10)
        c = nz.cell(row=i, column=2, value=help_txt)
        c.font = Font(name=FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")

    tips_row = 5 + len(COLUMNS) + 1
    tips = [
        "Astuce : mets 'publier' = non pour d'abord vérifier le rendu, puis oui pour uploader.",
        "traduire_vers : indique une chaîne d'une autre langue (ex en1) -> la vidéo est "
        "traduite et publiée aussi sur cette chaîne (nécessite une clé LLM).",
        "musique : dépose tes .mp3 libres de droits dans data/assets/music, puis écris "
        "le nom du fichier ici. 'aucune' = pas de musique. Vide = choix aléatoire.",
        "voix : liste complète via `edge-tts --list-voices`. Vide = voix de la chaîne.",
        "Avant d'uploader : autorise chaque chaîne avec `python main.py auth --channel frX`.",
    ]
    for j, t in enumerate(tips):
        cell = nz.cell(row=tips_row + j, column=1, value="• " + t)
        cell.font = Font(name=FONT, size=10)
        cell.alignment = Alignment(wrap_text=True)
        nz.merge_cells(start_row=tips_row + j, start_column=1,
                       end_row=tips_row + j, end_column=2)

    # ============ Feuille "Reference" ============
    rz = wb.create_sheet("Reference")
    for col, w in (("A", 34), ("B", 26), ("C", 22)):
        rz.column_dimensions[col].width = w
    rz.cell(row=1, column=1, value="Voix FR").font = Font(name=FONT, bold=True, size=12)
    for i, v in enumerate(VOICES_FR, start=2):
        rz.cell(row=i, column=1, value=v).font = Font(name=FONT, size=10)
    rz.cell(row=1, column=2, value="Voix EN").font = Font(name=FONT, bold=True, size=12)
    for i, v in enumerate(VOICES_EN, start=2):
        rz.cell(row=i, column=2, value=v).font = Font(name=FONT, size=10)

    base = len(VOICES_FR) + 3
    rz.cell(row=base, column=1, value="categoryId").font = Font(name=FONT, bold=True, size=12)
    rz.cell(row=base, column=2, value="Catégorie").font = Font(name=FONT, bold=True, size=12)
    for i, (cid, label) in enumerate(CATEGORIES, start=base + 1):
        rz.cell(row=i, column=1, value=cid).font = Font(name=FONT, size=10)
        rz.cell(row=i, column=2, value=label).font = Font(name=FONT, size=10)

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    import sys

    dest = sys.argv[1] if len(sys.argv) > 1 else "shorts_template.xlsx"
    print("Modèle créé :", make_template(dest))
