"""Mode batch : lit un fichier Excel et produit (+ uploade) un Short par ligne.

Une ligne = une vidéo. Colonnes reconnues (ordre libre, accents/casse ignorés) :

  channel            id de la chaîne cible (section `channels` du config)
  titre              titre YouTube (utilisé tel quel)
  script             texte lu par la voix off (si vide + clé LLM -> généré)
  topic              thème (sert au fallback / génération si pas de script)
  keywords           mots-clés (virgule) -> requête b-roll Pexels + tags
  hashtags           hashtags (virgule)
  cta                phrase de fin (call-to-action)
  voix               voix edge-tts (surcharge la voix de la chaîne)
  musique            nom de fichier dans data/assets/music (ou "aucune")
  fond               auto | pexels | local | color
  categorie          categoryId YouTube (nombre) — surcharge la chaîne
  traduire_vers      id d'une AUTRE chaîne : duplique la vidéo traduite dessus
  publier            oui/non — uploader ou seulement générer
  date_publication   date de programmation (ex 2026-08-01 09:00) ou vide

Tout est optionnel sauf de quoi produire un script (script OU clé LLM+topic).
Les erreurs d'une ligne n'arrêtent pas les autres : un récap est affiché.
"""
from __future__ import annotations

import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from .channels import Channel, get_channel
from .config import Config
from .logger import get_logger
from .pipeline import Pipeline

log = get_logger(__name__)

# --- mapping des en-têtes (forme normalisée -> clé interne) ---
_HEADER_ALIASES = {
    "channel": "channel", "chaine": "channel", "chaîne": "channel",
    "titre": "title", "title": "title",
    "script": "script", "texte": "script", "text": "script",
    "topic": "topic", "sujet": "topic", "theme": "topic", "thème": "topic",
    "keywords": "keywords", "motscles": "keywords", "mots-cles": "keywords",
    "mots_cles": "keywords", "tags": "keywords",
    "hashtags": "hashtags",
    "cta": "cta",
    "voix": "voice", "voice": "voice",
    "musique": "music", "music": "music",
    "fond": "background", "background": "background",
    "categorie": "category", "category": "category",
    "traduirevers": "translate_to", "traduire_vers": "translate_to",
    "traduire": "translate_to", "translateto": "translate_to",
    "publier": "publish", "publish": "publish", "upload": "publish",
    "datepublication": "publish_date", "date_publication": "publish_date",
    "date": "publish_date", "publishdate": "publish_date",
}

_TRUE = {"oui", "yes", "true", "vrai", "1", "x", "y", "o"}


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text or ""))
    text = "".join(c for c in text if not unicodedata.combining(c))
    return re.sub(r"[\s]+", "", text).strip().lower()


def _split_list(value: Any) -> list[str]:
    if value is None:
        return []
    parts = re.split(r"[;,|\n]+", str(value))
    return [p.strip() for p in parts if p.strip()]


def _to_bool(value: Any, default: bool) -> bool:
    if value is None or str(value).strip() == "":
        return default
    return str(value).strip().lower() in _TRUE


def _parse_date(value: Any) -> datetime | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M",
                "%d/%m/%Y %H:%M", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    log.warning("Date « %s » non reconnue -> ignorée.", s)
    return None


def _slug(text: str) -> str:
    text = _norm(text)
    text = re.sub(r"[^\w-]", "", text)
    return (text[:32] or "short")


def _read_rows(xlsx_path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Le module openpyxl est requis pour lire l'Excel "
            "(pip install openpyxl)."
        ) from exc

    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb["Shorts"] if "Shorts" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []

    headers = [_HEADER_ALIASES.get(_norm(h), None) for h in rows[0]]
    out: list[dict[str, Any]] = []
    for raw in rows[1:]:
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        rec: dict[str, Any] = {}
        for key, cell in zip(headers, raw):
            if key:
                rec[key] = cell
        out.append(rec)
    return out


def _build_topic(rec: dict[str, Any], index: int) -> dict[str, Any]:
    title = (rec.get("title") or "").strip() if rec.get("title") else None
    topic_theme = (rec.get("topic") or title or "").strip()
    topic: dict[str, Any] = {
        "id": f"{_slug(title or topic_theme or f'ligne{index}')}-r{index}",
        "topic": topic_theme,
        "keywords": _split_list(rec.get("keywords")),
        "hashtags": _split_list(rec.get("hashtags")),
    }
    if title:
        topic["title"] = title
    if rec.get("cta"):
        topic["cta"] = str(rec["cta"]).strip()
    if rec.get("script") and str(rec["script"]).strip():
        topic["script"] = str(rec["script"]).strip()
    return topic


def _line_overrides(rec: dict[str, Any]) -> dict[str, Any]:
    """Surcharges de config issues des colonnes voix/musique/fond."""
    ov: dict[str, Any] = {}
    if rec.get("voice"):
        ov["tts"] = {"voice": str(rec["voice"]).strip()}
    if rec.get("music") and str(rec["music"]).strip():
        ov["music"] = {"file": str(rec["music"]).strip()}
    if rec.get("background") and str(rec["background"]).strip():
        ov["background"] = {"source": str(rec["background"]).strip().lower()}
    return ov


def _merge(base: dict[str, Any], extra: dict[str, Any]) -> dict[str, Any]:
    """Deep-merge non destructif (extra prioritaire)."""
    out = {k: (dict(v) if isinstance(v, dict) else v) for k, v in base.items()}
    for key, value in extra.items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = _merge(out[key], value)
        else:
            out[key] = value
    return out


def _upload_settings(channel: Channel | None, rec: dict[str, Any]) -> dict[str, Any]:
    settings: dict[str, Any] = channel.upload_settings() if channel else {}
    if rec.get("category"):
        settings["category_id"] = str(rec["category"]).strip()
    when = _parse_date(rec.get("publish_date"))
    if when is not None:
        settings["publish_at"] = when
    return settings


def run_batch(cfg: Config, xlsx_path: str | Path, do_upload: bool = True,
              on_row=None) -> dict:
    """Traite tout le fichier Excel. Renvoie un récap {ok, failed, uploaded}.

    `on_row(index, total, titre)` est appelé au début de chaque ligne (pour
    remonter la progression à une interface, par ex. le dashboard).
    """
    xlsx_path = Path(xlsx_path)
    if not xlsx_path.exists():
        raise FileNotFoundError(f"Fichier introuvable : {xlsx_path}")

    records = _read_rows(xlsx_path)
    if not records:
        log.warning("Aucune ligne exploitable dans %s.", xlsx_path.name)
        return {"ok": 0, "failed": 0, "uploaded": 0, "total": 0}

    log.info("Batch : %d ligne(s) à traiter (%s).", len(records), xlsx_path.name)
    pipeline = Pipeline(cfg)
    ok = failed = uploaded = 0

    try:
        for i, rec in enumerate(records, start=1):
            title_preview = (rec.get("title") or rec.get("topic") or f"ligne {i}")
            if on_row:
                on_row(i, len(records), str(title_preview)[:60])
            log.info("──── [%d/%d] %s ────", i, len(records), str(title_preview)[:60])
            try:
                channel = get_channel(cfg, (rec.get("channel") or "").strip() or None)
                base_ov = channel.overrides() if channel else {}
                overrides = _merge(base_ov, _line_overrides(rec))
                topic = _build_topic(rec, i)

                record = pipeline.generate(topic, overrides=overrides)
                ok += 1
                print(f"  ✅ Généré : {Path(record['file_path']).name}")

                row_upload = do_upload and _to_bool(rec.get("publish"), default=True)
                if row_upload:
                    settings = _upload_settings(channel, rec)
                    yt = pipeline.upload(record, settings=settings)
                    uploaded += 1
                    print(f"  ⬆️  En ligne : https://youtu.be/{yt}")

                # ---- duplication traduite sur une autre chaîne ----
                target_id = (rec.get("translate_to") or "").strip()
                if target_id:
                    gen2, up2 = _duplicate_translated(
                        cfg, pipeline, rec, record, target_id, i, do_upload
                    )
                    ok += gen2
                    uploaded += up2

            except Exception as exc:  # noqa: BLE001
                failed += 1
                log.error("Ligne %d en échec : %s", i, exc)
                print(f"  ❌ Échec ligne {i} : {exc}")
    finally:
        pipeline.close()

    print(
        f"\n=== Batch terminé : {ok} générée(s), {uploaded} uploadée(s), "
        f"{failed} échec(s) sur {len(records)} ligne(s). ==="
    )
    return {"ok": ok, "failed": failed, "uploaded": uploaded, "total": len(records)}


def _duplicate_translated(
    cfg: Config,
    pipeline: Pipeline,
    rec: dict[str, Any],
    source_record: dict[str, Any],
    target_id: str,
    index: int,
    do_upload: bool,
) -> tuple[int, int]:
    """Traduit la vidéo source et la republie sur la chaîne `target_id`.

    Renvoie (nb_générées, nb_uploadées) pour cette version traduite.
    """
    from .generators.translate import has_llm, translate_bundle

    if not has_llm(cfg):
        print("  ⚠️  Traduction ignorée (aucune clé LLM dans .env).")
        return (0, 0)

    target = get_channel(cfg, target_id)
    if target is None:
        print(f"  ⚠️  Chaîne de traduction « {target_id} » inconnue -> ignorée.")
        return (0, 0)

    print(f"  🌍 Traduction vers « {target.name} » ({target.language})…")
    bundle = translate_bundle(
        source_record.get("script", ""),
        source_record["title"],
        source_record["description"],
        target.language,
        cfg,
    )

    topic = _build_topic(rec, index)
    topic["id"] = topic["id"] + f"-{target.language}"
    topic["title"] = bundle["title"] or source_record["title"]
    topic["description"] = bundle["description"] or source_record["description"]
    topic["script"] = bundle["script"]

    # Même musique/fond que la ligne, mais voix + langue de la chaîne cible
    # (la voix de la colonne, en FR, ne doit pas s'appliquer à la version traduite).
    line_ov = _line_overrides(rec)
    line_ov.pop("tts", None)
    overrides = _merge(line_ov, target.overrides())

    record = pipeline.generate(topic, overrides=overrides)
    print(f"  ✅ Version {target.language} : {Path(record['file_path']).name}")

    uploaded = 0
    if do_upload and _to_bool(rec.get("publish"), default=True):
        settings = _upload_settings(target, rec)
        yt = pipeline.upload(record, settings=settings)
        uploaded = 1
        print(f"  ⬆️  En ligne ({target.language}) : https://youtu.be/{yt}")
    return (1, uploaded)
